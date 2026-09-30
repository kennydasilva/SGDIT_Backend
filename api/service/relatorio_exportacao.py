from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from api.model.user import Admin
from api.service.relatorio_service import RelatorioService

ESTADO_LABEL = {
    "PENDENTE": "Pendente", "VALIDADA": "Validada", "APROVADA": "Aprovada",
    "REJEITADA": "Rejeitada", "ARQUIVADA": "Arquivada",
    "ENCAMINHADA": "Enviada ao posto", "EM_ATENDIMENTO": "Agente designado",
}
TIPO_LABEL = {
    "CONTRAMAO": "Contramão", "PARADO": "Veículo parado",
    "VELOCIDADE": "Excesso de velocidade", "ACIDENTE": "Acidente de viação",
}

# O PDF lista no máximo estas denúncias (o Excel lista todas).
MAX_LINHAS_PDF = 1000

COLUNAS_DETALHE = [
    "N.º", "Data", "Tipo", "Matrícula", "Estado", "Posto", "Localização",
    "Agente", "Código legal", "Ligada à", "Possível falsa",
]

AZUL = "1D4ED8"


def _fmt(valor, sufixo=""):
    return "—" if valor is None else f"{valor}{sufixo}"


def descrever_filtros(filtros):
    partes = []
    if filtros.get("data_inicio") or filtros.get("data_fim"):
        ini = filtros["data_inicio"].strftime("%d/%m/%Y") if filtros.get("data_inicio") else "início"
        fim = filtros["data_fim"].strftime("%d/%m/%Y") if filtros.get("data_fim") else "hoje"
        partes.append(f"Período: {ini} a {fim}")
    else:
        partes.append("Período: todo")
    if filtros.get("sem_posto"):
        partes.append("Posto: sem posto")
    elif filtros.get("admin_id"):
        posto = Admin.objects.filter(id=filtros["admin_id"]).values_list("posto", flat=True).first()
        partes.append(f"Posto: {posto or filtros['admin_id']}")
    else:
        partes.append("Posto: todos")
    if filtros.get("tipo"):
        partes.append(f"Tipo: {TIPO_LABEL.get(filtros['tipo'], filtros['tipo'])}")
    if filtros.get("estado"):
        partes.append(f"Estado: {ESTADO_LABEL.get(filtros['estado'], filtros['estado'])}")
    return " · ".join(partes)


def _linhas_detalhe(qs):
    for d in qs:
        analise = getattr(d, "resultado_analise", None)
        yield [
            d.id,
            timezone.localtime(d.data_registo).strftime("%d/%m/%Y %H:%M") if d.data_registo else "",
            TIPO_LABEL.get(d.tipo_infracao, d.tipo_infracao or ""),
            d.matricula or "",
            ESTADO_LABEL.get(d.estado, d.estado),
            d.admin_responsavel.posto if d.admin_responsavel_id else "Sem posto",
            d.localizacao or "",
            d.pt.utilizador.nome if d.pt_id else "",
            (analise.codigo_legal if analise and analise.codigo_legal else d.codigo_legal) or "",
            f"#{d.denuncia_principal_id}" if d.denuncia_principal_id else "",
            "Sim" if d.localizacao_contraditoria else "",
        ]


def _seccoes_resumo(resumo):
    """As mesmas tabelas para o Excel e para o PDF: (título, cabeçalho, linhas)."""
    q, a = resumo["qualidade"], resumo["acidentes"]
    return [
        ("Indicadores", ["Indicador", "Valor"], [
            ["Total de denúncias", resumo["total_denuncias"]],
            ["Taxa de resolução", f"{resumo['taxa_resolucao']}%"],
            ["Tempo médio de resposta", _fmt(resumo["tempo_medio_resposta_horas"], " h")],
        ]),
        ("Por estado", ["Estado", "Denúncias"],
         [[ESTADO_LABEL.get(k, k), v] for k, v in resumo["por_estado"].items() if v]),
        ("Por tipo de infração", ["Tipo", "Denúncias"],
         [[TIPO_LABEL.get(k, k), v] for k, v in resumo["por_tipo_infracao"].items()]),
        ("Por mês", ["Mês", "Denúncias"], [[m["mes"], m["total"]] for m in resumo["por_mes"]]),
        ("Por posto", ["Posto", "Denúncias", "Aprovadas", "Acidentes", "Taxa de resolução", "Tempo médio de resposta"],
         [[p["posto"], p["total"], p["aprovadas"], p["acidentes"], f"{p['taxa_resolucao']}%",
           _fmt(p["tempo_medio_resposta_horas"], " h")] for p in resumo["por_posto"]]),
        ("Desempenho dos agentes", ["Agente", "N.º", "Posto", "Decisões", "Aprovadas", "Arquivadas", "Tempo médio até decidir"],
         [[g["nome"], g["numero_agente"], g["posto"] or "", g["decisoes"], g["aprovadas"], g["arquivadas"],
           _fmt(g["tempo_medio_decisao_horas"], " h")] for g in resumo["agentes"]]),
        ("Acidentes", ["Indicador", "Valor"], [
            ["Acidentes (sem contar reportes repetidos)", a["total"]],
            ["Reportes juntados a um acidente já reportado", a["reportes_juntados"]],
            ["Fora de qualquer jurisdição", a["sem_posto"]],
            ["À espera de agente", a["a_aguardar_agente"]],
            ["Com agente designado", a["com_agente"]],
            ["Tempo médio até designar agente", _fmt(a["tempo_medio_ate_designar_min"], " min")],
        ]),
        ("Qualidade das denúncias", ["Indicador", "Valor"], [
            ["Rejeitadas pela análise automática", q["rejeitadas_ia"]],
            ["Vídeo ilegível (análise falhou)", q["video_ilegivel"]],
            ["Taxa de rejeição automática", f"{q['taxa_rejeicao_ia']}%"],
            ["Testemunhas (mesma infração, outro cidadão)", q["testemunhas"]],
            ["Vídeos semelhantes a outra denúncia", q["videos_semelhantes"]],
            ["Possíveis denúncias falsas (local contraditório)", q["possiveis_falsas"]],
        ]),
    ]


# ---- Excel -------------------------------------------------------------------

def gerar_excel(filtros):
    resumo = RelatorioService.obter_resumo(filtros)
    wb = Workbook()

    cabecalho_font = Font(bold=True, color="FFFFFF")
    cabecalho_fill = PatternFill("solid", fgColor=AZUL)

    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Relatório de denúncias — SGDIT"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([descrever_filtros(filtros)])
    ws.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])

    for titulo, cabecalho, linhas in _seccoes_resumo(resumo):
        ws.append([])
        ws.append([titulo])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(cabecalho)
        for col in range(1, len(cabecalho) + 1):
            c = ws.cell(row=ws.max_row, column=col)
            c.font, c.fill = cabecalho_font, cabecalho_fill
        for linha in linhas or [["Sem dados"]]:
            ws.append(linha)
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 26 if col == 1 else 18
    ws.column_dimensions["A"].width = 48

    wd = wb.create_sheet("Denúncias")
    wd.append(COLUNAS_DETALHE)
    for col in range(1, len(COLUNAS_DETALHE) + 1):
        c = wd.cell(row=1, column=col)
        c.font, c.fill = cabecalho_font, cabecalho_fill
    for linha in _linhas_detalhe(RelatorioService.listar_detalhe(filtros)):
        wd.append(linha)
    larguras = [8, 17, 22, 12, 18, 28, 45, 24, 26, 10, 12]
    for i, largura in enumerate(larguras, start=1):
        wd.column_dimensions[get_column_letter(i)].width = largura
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = wd.dimensions

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ---- PDF ---------------------------------------------------------------------

def _tabela_pdf(cabecalho, linhas, larguras=None, tamanho=8):
    estilo_celula = ParagraphStyle("celula", fontSize=tamanho, leading=tamanho + 2)
    # Cabeçalho também em Paragraph, para quebrar linha em colunas estreitas.
    estilo_cabecalho = ParagraphStyle(
        "cabecalho", parent=estilo_celula, textColor=colors.white, fontName="Helvetica-Bold"
    )
    dados = [[Paragraph(str(c), estilo_cabecalho) for c in cabecalho]] + [
        [Paragraph(str(v), estilo_celula) for v in linha] for linha in (linhas or [["Sem dados"]])
    ]
    t = Table(dados, colWidths=larguras, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + AZUL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), tamanho),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return t


def gerar_pdf(filtros):
    resumo = RelatorioService.obter_resumo(filtros)
    saida = BytesIO()
    doc = SimpleDocTemplate(
        saida, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title="Relatório de denúncias — SGDIT",
    )
    estilos = getSampleStyleSheet()
    h1, h2, normal = estilos["Title"], estilos["Heading2"], estilos["Normal"]
    subtitulo = ParagraphStyle("sub", parent=normal, textColor=colors.HexColor("#6B7280"), fontSize=9)

    corpo = [
        Paragraph("Relatório de denúncias — SGDIT", h1),
        Paragraph(descrever_filtros(filtros), subtitulo),
        Paragraph(f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}", subtitulo),
        Spacer(1, 0.4 * cm),
    ]
    for titulo, cabecalho, linhas in _seccoes_resumo(resumo):
        # Título e tabela na mesma página (o título nunca fica sozinho no fundo).
        corpo += [KeepTogether([Paragraph(titulo, h2), _tabela_pdf(cabecalho, linhas)]), Spacer(1, 0.3 * cm)]

    qs = RelatorioService.listar_detalhe(filtros)
    total = qs.count()
    corpo.append(Paragraph("Lista de denúncias", h2))
    if total > MAX_LINHAS_PDF:
        corpo.append(Paragraph(
            f"Mostradas as {MAX_LINHAS_PDF} mais recentes de {total}. A lista completa está na exportação Excel.",
            subtitulo,
        ))
    # Soma = 26,7 cm: a largura útil de um A4 horizontal com estas margens.
    larguras = [1.2, 2.6, 2.8, 2.0, 2.4, 3.0, 4.6, 2.6, 2.6, 1.4, 1.5]
    corpo.append(_tabela_pdf(
        COLUNAS_DETALHE,
        list(_linhas_detalhe(qs[:MAX_LINHAS_PDF])),
        larguras=[l * cm for l in larguras],
        tamanho=7,
    ))

    def rodape(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#9CA3AF"))
        canvas.drawRightString(landscape(A4)[0] - 1.5 * cm, 0.8 * cm, f"SGDIT · página {doc_.page}")
        canvas.restoreState()

    doc.build(corpo, onFirstPage=rodape, onLaterPages=rodape)
    return saida.getvalue()
